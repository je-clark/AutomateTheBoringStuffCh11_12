import openpyxl

fn = 'example.xlsx'

wb = openpyxl.load_workbook(fn)
sheet = wb.active
refObj1 = openpyxl.chart.Reference(sheet, min_col=1, min_row=2, max_col=1, max_row=6)
seriesObj1 = openpyxl.chart.Series(refObj1, title=sheet['A1'].value)

refObj2 = openpyxl.chart.Reference(sheet, min_col=2, min_row=2, max_col=2, max_row=6)
seriesObj2 = openpyxl.chart.Series(refObj2, title=sheet['B1'].value)

chartObj = openpyxl.chart.BarChart()
chartObj.append(seriesObj1)
chartObj.append(seriesObj2)
wb.create_sheet(index=1, title="chart").add_chart(chartObj)

wb.save('example_edited.xlsx')